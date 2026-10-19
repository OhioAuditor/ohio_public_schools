
import json
import logging
import os
import re
import tempfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import PurePosixPath
from urllib.parse import unquote_plus

import boto3
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook


# ============================================================
# AWS CONFIGURATION
# ============================================================

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

AWS_REGION = os.getenv("AWS_REGION", "us-east-2")
RAW_BUCKET = os.getenv("RAW_BUCKET", "schools-raw").strip()
DEST_BUCKET = os.getenv(
    "DEST_BUCKET", "schools-transformed"
).strip()

SOURCE_PREFIX = os.getenv(
    "CANTON_PUBLIC_SOURCE_PREFIX",
    "data-from-clients/canton_public/",
).strip("/") + "/"

CHUNK_ROWS = max(
    1000,
    int(os.getenv("ERNREG_CHUNK_ROWS", "10000")),
)


# ============================================================
# REPORT COLUMNS
# ============================================================

PAYROLL_COLUMNS = [
    "id",
    "name",
    "ck_id",
    "ck_no",
    "ck_date",
    "ck_amt",
    "ck_status",
    "subsystem",
    "job_no",
]

CHECK_COLUMNS = [
    "check_id",
    "check_no",
    "check_date",
    "check_status",
    "vendor_id",
    "vendor_name",
    "po_no",
    "invoice_no",
    "invoice_date",
    "description",
    "amount",
    "org_key",
    "fund",
    "scc",
    "func",
    "obj",
    "subj",
    "opu",
    "il",
    "job",
]

CASH_COLUMNS = [
    "fund",
    "scc",
    "fndscc_title",
    "begin_balance",
    "fytd_receipts",
    "fytd_expenditures",
    "end_balance",
    "encumbrances",
    "unencumbered_balance",
]

ERNREG_COLUMNS = [
    "id",
    "name",
    "period",
    "chk",
    "chk_dt",
    "tc_batch",
    "payclass",
    "position",
    "year",
    "step",
    "cdh",
    "cdh_desc",
    "amount",
    "gl_key",
    "key_desc",
    "fund",
    "scc",
    "func",
    "obj",
    "subj",
    "opu",
    "il",
    "job",
    "numcd",
]

ERNREG_DEDUCTION_COLUMNS = [
    "id",
    "name",
    "period",
    "chk",
    "chk_dt",
    "tc_batch",
    "payclass",
    "position",
    "year",
    "step",
    "cdh",
    "cdh_desc",
    "amount",
    "gl_key",
    "gl_obj",
    "numcd",
]


BUDLED_COLUMNS = [
    "key", "key_title", "fnd", "scc", "func", "obj",
    "subj", "opu", "il", "job", "inv_no", "ref_po",
    "pe_id", "pe_name", "ck_id", "ck_no", "desc",
    "trans_dt", "amount",
]


REVLED_COLUMNS = [
    "org_key", "fund", "scc", "rcpt", "subj", "opu",
    "il", "job", "pri_ref", "ref_po", "vend_id", "ck_no",
    "description", "s_ref_dt", "amount",
]


# ============================================================
# DESTINATION CONFIGURATION
# ============================================================

def destination(env, default):
    return os.getenv(
        env,
        "data_from_clients/" + default + "/",
    ).strip("/")


# ============================================================
# REPORT DEFINITIONS
# ============================================================

REPORTS = {
    "payroll_check_reg": {
        "phrase": "payroll check register",
        "columns": PAYROLL_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_PAYROLL_CHECK_REG_DEST_PREFIX",
            "canton_public_payroll_check_reg",
        ),
        "dates": ["ck_date"],
        "amounts": ["ck_amt"],
    },

    "check_register": {
        "phrase": "check register",
        "columns": CHECK_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_CHECK_REGISTER_DEST_PREFIX",
            "canton_public_check_register",
        ),
        "dates": ["check_date", "invoice_date"],
        "amounts": ["amount"],
    },

    "cash_fund_bal": {
        "phrase": "finsum",
        "columns": CASH_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_CASH_FUND_BAL_DEST_PREFIX",
            "canton_public_cash_fund_bal",
        ),
        "dates": [],
        "amounts": [
            "begin_balance",
            "fytd_receipts",
            "fytd_expenditures",
            "end_balance",
            "encumbrances",
            "unencumbered_balance",
        ],
    },

    "ernreg_wages_benefits": {
        "phrase": "ernreg wages and benefits",
        "columns": ERNREG_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_ERNREG_WAGES_BENEFITS_DEST_PREFIX",
            "canton_public_ernreg_wages_benefits",
        ),
        "dates": ["chk_dt"],
        "amounts": ["amount"],
    },

    "ernreg_deduction": {
        "phrase": "ernreg deductions",
        "columns": ERNREG_DEDUCTION_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_ERNREG_DEDUCTION_DEST_PREFIX",
            "canton_public_ernreg_deduction",
        ),
        "dates": ["chk_dt"],
        "amounts": ["amount"],
    },
    "budled": {
        "phrase": "budled",
        "columns": BUDLED_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_BUDLED_DEST_PREFIX",
            "canton_public_budled",
        ),
        "dates": ["trans_dt"],
        "amounts": ["amount"],
        "header_row": 4,
    },

    "revled": {
        "phrase": "revled",
        "columns": REVLED_COLUMNS,
        "destination": destination(
            "CANTON_PUBLIC_REVLED_DEST_PREFIX",
            "canton_public_revled",
        ),
        "dates": ["s_ref_dt"],
        "amounts": ["amount"],
        "header_row": 4,
    },
}


# ============================================================
# COLUMN NORMALIZATION
# ============================================================

def normalize_col_name(name):
    text = str(name).strip().lower()
    text = text.replace("#", "number")

    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text)

    return text.strip("_")


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(value):
    if value is None or pd.isna(value):
        return None

    return str(value).strip()


# ============================================================
# CURRENCY CONVERSION
# ============================================================

def currency(value):
    """
    Convert currency to exact cents.

    Examples:
        1234.567 -> 1234.57
        (500.25) -> -500.25
        $1,250.00 -> 1250.00
    """

    if value is None or pd.isna(value):
        return None

    cleaned = str(value).strip()

    if not cleaned:
        return None

    if cleaned.startswith("(") and cleaned.endswith(")"):
        cleaned = "-" + cleaned[1:-1]

    cleaned = re.sub(r"[$,%\s,]", "", cleaned)

    try:
        return Decimal(cleaned).quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        )

    except InvalidOperation as exc:
        raise ValueError(
            f"Invalid currency amount: {value!r}"
        ) from exc


# ============================================================
# REPORT IDENTIFICATION
# ============================================================

def report_type_for_key(key):
    filename = PurePosixPath(key).stem.lower()

    filename = re.sub(
        r"[^a-z0-9]+",
        " ",
        filename,
    )

    filename = re.sub(
        r"\s+",
        " ",
        filename,
    ).strip()

    # Payroll must be checked before Check Register.
    report_order = [
        "payroll_check_reg",
        "check_register",
        "cash_fund_bal",
        "ernreg_wages_benefits",
        "ernreg_deduction",
        "budled",
        "revled",
    ]

    for report_type in report_order:
        phrase = REPORTS[report_type]["phrase"]

        if phrase in filename:
            return report_type

    return None


# ============================================================
# FISCAL YEAR DETECTION
# ============================================================

def fiscal_year_for_key(key):
    filename = PurePosixPath(key).name

    match = re.search(
        r"(?i)(?<![a-z0-9])FY[\s_-]*(20\d{2}|\d{2})(?!\d)",
        filename,
    )

    if match:
        year = match.group(1)

        if len(year) == 2:
            return "20" + year

        return year

    match = re.search(
        r"(?<!\d)(20\d{2})(?!\d)",
        filename,
    )

    if match:
        return match.group(1)

    raise ValueError(
        f"Could not determine fiscal year: {filename}"
    )


# ============================================================
# PARQUET SCHEMA
# ============================================================

def schema_for(report_type):
    config = REPORTS[report_type]

    fields = []

    for col in config["columns"]:
        if col in config["amounts"]:
            dtype = pa.decimal128(20, 2)

        elif col in config["dates"]:
            dtype = pa.date32()

        else:
            dtype = pa.string()

        fields.append(pa.field(col, dtype))

    fields.append(
        pa.field("fiscal_year", pa.string())
    )

    return pa.schema(fields)


# ============================================================
# DATAFRAME TRANSFORMATION
# ============================================================

def transform_df(df, report_type, fiscal_year):
    config = REPORTS[report_type]

    expected = config["columns"]
    actual = list(df.columns)

    if actual != expected:
        raise ValueError(
            f"Unexpected {report_type} columns: "
            f"{actual}; expected {expected}"
        )

    df = df.dropna(how="all").copy()

    for col in expected:
        if col in config["amounts"]:
            df[col] = df[col].map(currency)

        elif col in config["dates"]:
            df[col] = pd.to_datetime(
                df[col],
                errors="coerce",
            ).dt.date

        else:
            df[col] = df[col].map(
                clean_text
            ).astype("string")

    df["fiscal_year"] = fiscal_year

    return df


# ============================================================
# SMALL EXCEL REPORT CONVERSION
# ============================================================

def convert_small_report(
    content,
    report_type,
    fiscal_year,
    output_path,
):
    columns = REPORTS[report_type]["columns"]

    df = pd.read_excel(
        BytesIO(content),
        sheet_name=0,
        dtype=object,
        engine="openpyxl",
        usecols=range(len(columns)),
    )

    df.columns = [
        normalize_col_name(col)
        for col in df.columns
    ]

    df = transform_df(
        df,
        report_type,
        fiscal_year,
    )

    table = pa.Table.from_pandas(
        df,
        schema=schema_for(report_type),
        preserve_index=False,
    )

    pq.write_table(
        table,
        output_path,
        compression="snappy",
    )

    return len(df)


# ============================================================
# LARGE ERNREG REPORT CONVERSION
# ============================================================

def convert_large_report(source_file, report_type, fiscal_year, output_path):
    """Stream ERNREG, BUDLED, and REVLED in fixed-size Parquet batches."""
    workbook = load_workbook(source_file, read_only=True, data_only=True)
    writer = None
    total_rows = 0
    try:
        worksheet = workbook.active
        expected = REPORTS[report_type]["columns"]
        header_row = REPORTS[report_type].get("header_row", 1)
        rows = worksheet.iter_rows(min_row=header_row, values_only=True)
        headers = next(rows)
        actual = [normalize_col_name(v) for v in headers[:len(expected)]]
        if actual != expected:
            raise ValueError(
                f"Unexpected {report_type} columns: {actual}; expected {expected}"
            )

        schema = schema_for(report_type)
        batch = []

        def flush(items, current_writer):
            frame = pd.DataFrame(items, columns=expected)
            frame = transform_df(frame, report_type, fiscal_year)
            table = pa.Table.from_pandas(
                frame, schema=schema, preserve_index=False
            )
            if current_writer is None:
                current_writer = pq.ParquetWriter(
                    output_path, schema, compression="snappy"
                )
            current_writer.write_table(table)
            return current_writer, len(frame)

        for row in rows:
            values = tuple(row[:len(expected)])
            if all(v is None or (isinstance(v, str) and not v.strip())
                   for v in values):
                continue
            batch.append(values)
            if len(batch) >= CHUNK_ROWS:
                writer, count = flush(batch, writer)
                total_rows += count
                logger.info("%s: processed %s rows", report_type, total_rows)
                batch = []

        if batch:
            writer, count = flush(batch, writer)
            total_rows += count

        if writer is None:
            writer = pq.ParquetWriter(
                output_path, schema, compression="snappy"
            )
        logger.info("Completed %s: %s rows", report_type, total_rows)
        return total_rows
    finally:
        if writer is not None:
            writer.close()
        workbook.close()


# ============================================================
# S3 CONVERSION
# ============================================================

def convert_s3_object(bucket, key, s3=None):
    if bucket != RAW_BUCKET:
        raise ValueError(
            f"Unexpected source bucket: {bucket}"
        )

    if not key.startswith(SOURCE_PREFIX):
        logger.info(
            "Skipping unrelated source prefix: %s",
            key,
        )
        return None

    filename = PurePosixPath(key).name

    if (
        not filename.lower().endswith(".xlsx")
        or filename.startswith("~$")
    ):
        logger.info(
            "Skipping non-Excel file: %s",
            key,
        )
        return None

    report_type = report_type_for_key(key)

    if report_type is None:
        logger.info(
            "Skipping unrecognized Canton Public file: %s",
            key,
        )
        return None

    fiscal_year = fiscal_year_for_key(key)

    if s3 is None:
        s3 = boto3.client(
            "s3",
            region_name=AWS_REGION,
        )

    config = REPORTS[report_type]

    logger.info(
        "Processing %s: s3://%s/%s",
        report_type,
        bucket,
        key,
    )

    with tempfile.TemporaryDirectory() as temp_dir:
        output_path = os.path.join(
            temp_dir,
            "output.parquet",
        )

        if report_type in (
            "ernreg_wages_benefits",
            "ernreg_deduction",
            "budled",
            "revled",
        ):
            # Large files are downloaded to disk.
            input_path = os.path.join(
                temp_dir,
                "source.xlsx",
            )

            s3.download_file(
                bucket,
                key,
                input_path,
            )

            with open(input_path, "rb") as handle:
                rows = convert_large_report(
                    handle,
                    report_type,
                    fiscal_year,
                    output_path,
                )

        else:
            # Smaller files are processed in memory.
            response = s3.get_object(
                Bucket=bucket,
                Key=key,
            )

            content = response["Body"].read()

            rows = convert_small_report(
                content,
                report_type,
                fiscal_year,
                output_path,
            )

        stem = PurePosixPath(key).stem

        dest_key = (
            f"{config['destination']}/{stem}.parquet"
        )

        s3.upload_file(
            output_path,
            DEST_BUCKET,
            dest_key,
        )

    logger.info(
        "Converted %s rows to s3://%s/%s",
        rows,
        DEST_BUCKET,
        dest_key,
    )

    return {
        "source": f"s3://{bucket}/{key}",
        "destination": (
            f"s3://{DEST_BUCKET}/{dest_key}"
        ),
        "report_type": report_type,
        "table": f"canton_public_{report_type}",
        "rows": rows,
        "fiscal_year": fiscal_year,
    }


# ============================================================
# S3 / SQS / SNS EVENT HANDLING
# ============================================================

def iter_s3_records(event):
    if isinstance(event, str):
        event = json.loads(event)

    if "Records" not in event:
        if "bucket" in event and "key" in event:
            yield event["bucket"], event["key"]

        return

    for record in event["Records"]:
        if "s3" in record:
            yield (
                record["s3"]["bucket"]["name"],
                unquote_plus(
                    record["s3"]["object"]["key"]
                ),
            )

        elif "body" in record:
            yield from iter_s3_records(
                record["body"]
            )

        elif "Sns" in record:
            yield from iter_s3_records(
                record["Sns"]["Message"]
            )


# ============================================================
# LAMBDA HANDLER
# ============================================================

def lambda_handler(event, context):
    results = []

    for bucket, key in iter_s3_records(event):
        result = convert_s3_object(
            bucket,
            key,
        )

        if result:
            results.append(result)

    logger.info(
        "Completed %s Canton Public conversions",
        len(results),
    )

    return {
        "statusCode": 200,
        "body": json.dumps(results),
    }
